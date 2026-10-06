"""Official transport data: station boardings, census commuting, bus and rail.

Requires openpyxl, as do the existing Excel collectors. No API key required.
Cached originals and their SHA-256 are retained; unknown layouts fail closed.
"""
import argparse
import hashlib
import html
import io
import json
import re
import time
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

try:
    from core.source_run import SourceRun
except ModuleNotFoundError:
    from collector.core.source_run import SourceRun

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / 'data/raw/transport'
OUT = ROOT / 'public/data/transport'
STATION_PAGE = 'https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-S12-2024.html'
CENSUS_ID = '000032214690'
BUS_LIST = 'https://www.e-stat.go.jp/stat-search/files?layout=dataset&toukei=00600330&tstat=000001078083&cycle=8'
RAIL_LIST = 'https://www.e-stat.go.jp/stat-search/files?layout=dataset&toukei=00600350&tstat=000001011026&cycle=1'
MODE_CODES = ['1', '21', '22', '23', '24', '25', '26', '27', '28', '31', '32', '33', '34', '35', '36', '4', '5']
PREFS = ['北海道','青森県','岩手県','宮城県','秋田県','山形県','福島県','茨城県','栃木県','群馬県','埼玉県','千葉県','東京都','神奈川県','新潟県','富山県','石川県','福井県','山梨県','長野県','岐阜県','静岡県','愛知県','三重県','滋賀県','京都府','大阪府','兵庫県','奈良県','和歌山県','鳥取県','島根県','岡山県','広島県','山口県','徳島県','香川県','愛媛県','高知県','福岡県','佐賀県','長崎県','熊本県','大分県','宮崎県','鹿児島県','沖縄県']


def fetch(url):
    request = urllib.request.Request(url, headers={'User-Agent': 'DATLUME public-data collector (https://datlume.com/about-data/)'})
    for attempt in range(3):
        try:
            with urllib.request.urlopen(request, timeout=40) as response:
                return response.read()
        except OSError:
            if attempt == 2:
                raise
            time.sleep(2 ** attempt)


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix('.tmp')
    tmp.write_text(json.dumps(value, ensure_ascii=False, separators=(',', ':')) + '\n')
    tmp.replace(path)


def acquire(name, url, offline):
    path = RAW / name
    if not offline:
        payload = fetch(url)
        # Archive by content hash before replacing the current working original.
        archive = RAW / 'archive' / (hashlib.sha256(payload).hexdigest() + path.suffix)
        archive.parent.mkdir(parents=True, exist_ok=True)
        if not archive.exists():
            archive.write_bytes(payload)
        path.write_bytes(payload)
    payload = path.read_bytes()
    return payload, {'file': name, 'url': url, 'sha256': hashlib.sha256(payload).hexdigest(), 'bytes': len(payload)}


def discover(listing, title, kind, cache, offline):
    if offline:
        return json.loads((RAW / cache).read_text())
    page = fetch(listing).decode('utf-8')
    for article in re.findall(r'<article\b.*?</article>', page, re.S):
        text = ' '.join(html.unescape(re.sub('<[^>]+>', ' ', article)).split())
        if title not in text:
            continue
        link = re.search(r'href="(/stat-search/file-download\?statInfId=(\d+)&(?:amp;)?fileKind=' + str(kind) + r')"', article)
        if link:
            record = {'url': 'https://www.e-stat.go.jp' + html.unescape(link[1]), 'id': link[2], 'description': text}
            write(RAW / cache, record)
            return record
    raise ValueError('Official table not found: ' + title)


def number(value, dash_zero=False):
    if value is None or value in ('', '…', '...', 'x', 'X'):
        return None
    if value in ('-', '－', '—'):
        return 0 if dash_zero else None
    result = int(value)
    if result < 0 or float(value) != result:
        raise ValueError('Unexpected count: ' + str(value))
    return result


def station_rows(payload):
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        names = [n for n in archive.namelist() if '/UTF-8/' in n and n.endswith('NumberOfPassengers.geojson')]
        if len(names) != 1:
            raise ValueError('Station GeoJSON layout changed')
        features = json.loads(archive.read(names[0]))['features']
    rows = []
    for feature_index, f in enumerate(features):
        p = f['properties']
        coords = f['geometry']['coordinates']
        if f['geometry']['type'] != 'LineString' or not coords:
            raise ValueError('Station geometry changed')
        values, states, notes = [], [], {}
        for offset, year in enumerate(range(2011, 2025)):
            base = 6 + 4 * offset
            duplicate, exists = int(p[f'S12_{base:03}']), int(p[f'S12_{base+1:03}'])
            known_gap = (p['S12_001c'], p['S12_001'], p['S12_002'], p['S12_003']) == ('010175', '大聖寺', 'IRいしかわ鉄道', 'IRいしかわ鉄道線') and year in range(2012, 2019) and (duplicate, exists, p[f'S12_{base+3:03}']) == (0, 0, 0)
            if not known_gap and (duplicate not in (1, 2, 3) or exists not in (1, 2, 3, 4)):
                raise ValueError('Unknown station status code')
            state = 'unclassified' if known_gap else 'other-line' if duplicate == 2 else 'absent' if duplicate == 3 or exists == 4 else {1:'ok', 2:'missing', 3:'private'}[exists]
            value = number(p[f'S12_{base+3:03}']) if state == 'ok' else None
            if state == 'ok' and value is None:
                raise ValueError('Station marked available without value')
            values.append(value)
            states.append(state)
            if p.get(f'S12_{base+2:03}'):
                notes[str(year)] = p[f'S12_{base+2:03}']
        identity = [p[k] for k in ('S12_001c','S12_001','S12_002','S12_003')]
        # Published station codes repeat for split geometry and use 000000 for
        # some historical stations. Match the full source identity, never name alone.
        sid = 's' + hashlib.sha256(json.dumps(identity, ensure_ascii=False).encode()).hexdigest()[:16]
        rows.append({'id': sid, 'code': p['S12_001c'], 'sourceFeatures': [feature_index], 'group': p['S12_001g'], 'name': p['S12_001'], 'operator': p['S12_002'], 'line': p['S12_003'],
                     'lon': round(sum(c[0] for c in coords)/len(coords), 6), 'lat': round(sum(c[1] for c in coords)/len(coords), 6),
                     'values': values, 'states': states, 'notes': notes})
    grouped = {}
    for r in rows:
        if r['id'] not in grouped:
            grouped[r['id']] = r
            continue
        previous = grouped[r['id']]
        previous['sourceFeatures'].extend(r['sourceFeatures'])
        for i in range(14):
            values = {v for v in (previous['values'][i], r['values'][i]) if v is not None}
            if len(values) > 1:
                raise ValueError('Conflicting repeated station values: ' + r['id'])
            if values:
                previous['values'][i] = values.pop()
                previous['states'][i] = 'ok'
            elif previous['states'][i] == 'other-line' and r['states'][i] != 'other-line':
                previous['states'][i] = r['states'][i]
        for year, note in r['notes'].items():
            if year in previous['notes'] and note != previous['notes'][year]:
                previous['notes'][year] += ' / ' + note
            else:
                previous['notes'][year] = note
    rows = list(grouped.values())
    if len(rows) < 10000:
        raise ValueError('Station coverage/identifier regression')
    return rows


def commute_rows(payload):
    sheet = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)['e17_02']
    rows = list(sheet.values)
    if rows[7][4] != '0_常住地による人口' or '１７－２表' not in rows[1][0]:
        raise ValueError('Census columns changed')
    grouped, labels = {}, {}
    for row in rows[10:]:
        if not row[2] or not row[3]:
            continue
        code, name = row[2].split('_', 1)
        mode, label = row[3].split('_', 1)
        if not re.fullmatch(r'\d{5}', code):
            raise ValueError('Census area identifier changed')
        rec = grouped.setdefault(code, {'code': code, 'name': name, 'prefecture': row[1].split('_', 1)[1], 'kind': str(row[0]), 'all': {}})
        if mode in rec['all']:
            raise ValueError('Duplicate census cell')
        # Census '-' means no applicable persons, not an unknown observation.
        rec['all'][mode] = number(row[4], dash_zero=True)
        labels[mode] = label
    result = []
    for r in grouped.values():
        v = r.pop('all')
        r['total'] = v['0']
        r['values'] = [v[c] for c in MODE_CODES]
        if None in r['values'] or sum(r['values']) != r['total']:
            raise ValueError('Census mutually exclusive categories do not sum: ' + r['code'])
        result.append(r)
    if len(result) < 1900:
        raise ValueError('Census area coverage regression')
    return result, [{'code': c, 'label': labels[c]} for c in MODE_CODES]


def bus_rows(payload):
    rows = list(openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True).worksheets[0].values)
    if '都道府県別（支局別）・車種別輸送人員' not in rows[0][0] or rows[2][9] != '単位：千人' or rows[6][3] != '計':
        raise ValueError('Bus columns or unit changed')
    year = int(re.search(r'(20\d{2})年度', rows[2][0])[1])
    records, hokkaido = [], []
    # Published values are rounded to thousands. Preserve that resolution.
    for row in rows[8:]:
        if row[0] == '全国計':
            records.append({'code': '00', 'name': '全国', 'thousands': number(row[3])})
        elif row[1] in ('札幌', '函館', '旭川', '室蘭', '釧路', '帯広', '北見'):
            hokkaido.append({'name': row[1], 'thousands': number(row[3])})
        elif row[1]:
            candidates = [i for i,p in enumerate(PREFS, 1) if p.removesuffix('都').removesuffix('府').removesuffix('県') == row[1]]
            if len(candidates) != 1:
                raise ValueError('Unknown bus prefecture: ' + row[1])
            records.append({'code': str(candidates[0]).zfill(2), 'name': PREFS[candidates[0]-1], 'thousands': number(row[3])})
    if len(hokkaido) != 7 or len(records) != 47:
        raise ValueError('Bus prefecture coverage changed')
    records.append({'code': '01', 'name': '北海道', 'thousands': sum(r['thousands'] for r in hokkaido), 'branches': hokkaido})
    records.sort(key=lambda r:r['code'])
    # Independent rounding of branch totals can differ slightly from national.
    if abs(sum(r['thousands'] for r in records[1:]) - records[0]['thousands']) > 54:
        raise ValueError('Bus national/branch sum mismatch')
    return year, records


def rail_rows(payload):
    sheet = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)['旅客数量（機械判読用）']
    rows = list(sheet.values)
    if rows[4][1] != '千人' or rows[3][1] != '全国計' or 'ＪＲ旅客会社' not in rows[3][3]:
        raise ValueError('Rail columns changed')
    annual, monthly = [], []
    for row in rows[5:]:
        label = str(row[0] or '')
        y = re.fullmatch(r'(20\d{2})年度', label)
        m = re.fullmatch(r'(20\d{2})年(\d{1,2})月', label)
        if not y and not m:
            continue
        values = [number(row[i]) for i in (1,3,5)]
        if any(v is None for v in values) or abs(values[0]-values[1]-values[2]) > 1:
            raise ValueError('Rail national/JR/private sum mismatch')
        (annual if y else monthly).append({'period': y[1] if y else m[1]+'-'+m[2].zfill(2), 'thousands': values})
    if len(annual) < 5 or len(monthly) < 13:
        raise ValueError('Rail history coverage changed')
    return annual, monthly


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--offline', action='store_true', help='Rebuild from acquired originals; no network')
    args = parser.parse_args()
    RAW.mkdir(parents=True, exist_ok=True)
    with SourceRun('public_transport', '国土交通省・総務省 公共交通') as run:
        station_url = 'https://nlftp.mlit.go.jp/ksj/gml/data/S12/S12-25/S12-25_GML.zip'
        # S12 schema/years are pinned. A new schema is deliberately reviewed first.
        station_raw, station_source = acquire('S12-25_GML.zip', station_url, args.offline)
        commute_raw, commute_source = acquire('commute.xlsx', 'https://www.e-stat.go.jp/stat-search/file-download?statInfId='+CENSUS_ID+'&fileKind=0', args.offline)
        bus = discover(BUS_LIST, '都道府県別（支局別）・車種別輸送人員', 4, 'bus-table.json', args.offline)
        rail = discover(RAIL_LIST, '鉄・軌道旅客輸送総括表', 0, 'rail-table.json', args.offline)
        bus_raw, bus_source = acquire('bus.xlsx', bus['url'], args.offline)
        rail_raw, rail_source = acquire('rail.xlsx', rail['url'], args.offline)
        stations = station_rows(station_raw)
        areas, modes = commute_rows(commute_raw)
        bus_year, buses = bus_rows(bus_raw)
        annual, monthly = rail_rows(rail_raw)
        stamp = datetime.now(timezone.utc).isoformat()
        sources = {'stations': {**station_source, 'page': STATION_PAGE, 'license': 'CC BY 4.0'},
                   'commute': {**commute_source, 'page': 'https://www.e-stat.go.jp/stat-search/files?stat_infid='+CENSUS_ID},
                   'bus': {**bus_source, 'page': 'https://www.e-stat.go.jp/stat-search/files?stat_infid='+bus['id']},
                   'rail': {**rail_source, 'page': 'https://www.e-stat.go.jp/stat-search/files?stat_infid='+rail['id']}}
        # All sources are parsed and validated before touching any public file.
        stations.sort(key=lambda r: (-(r['values'][-1] or -1), r['id']))
        write(OUT/'stations.json', {'schemaVersion':1, 'years':list(range(2011,2025)), 'records':stations})
        write(OUT/'commute.json', {'year':2020, 'modes':modes, 'records':areas})
        write(OUT/'usage.json', {'bus': {'year':bus_year, 'records':buses}, 'rail': {'annual':annual, 'monthly':monthly}})
        write(OUT/'index.json', {'schemaVersion':1, 'generatedAt':stamp, 'stationCount':len(stations), 'stationAvailableCount':sum(r['values'][-1] is not None for r in stations),
                              'commuteAreaCount':len(areas), 'stationYears':[2011,2024], 'commuteYear':2020, 'busYear':bus_year, 'railLatestMonth':monthly[-1]['period'], 'sources':sources})
        write(RAW/'manifest.json', {'generatedAt':stamp, 'sources':sources})
        run.set_metrics(records=len(stations)+len(areas)+len(buses)+len(annual)+len(monthly), stations=len(stations), areas=len(areas), latestMonth=monthly[-1]['period'])
        print(json.dumps({'stations':len(stations), 'areas':len(areas), 'busYear':bus_year, 'railLatestMonth':monthly[-1]['period']}, ensure_ascii=False))


if __name__ == '__main__':
    main()
