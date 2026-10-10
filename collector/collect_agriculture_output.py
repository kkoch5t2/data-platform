#!/usr/bin/env python3
"""Publish MAFF's corrected 2024 municipal agricultural output detail table."""
from __future__ import annotations

import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook

try:
    from core.source_run import SourceRun
except ModuleNotFoundError:
    from collector.core.source_run import SourceRun

ROOT = Path(__file__).resolve().parents[1]
SOURCE = "https://www.maff.go.jp/j/tokei/kouhyou/sityoson_sansyutu/attach/xls/index-25.xlsx"
RAW = ROOT / "data/raw/agriculture-output/2024-detail.xlsx"
PUBLIC = ROOT / "public/data/agriculture-output.json"
HEADERS = (
    "年次", "都道府県コード", "市町村コード", "全国農業地域", "都道府県", "市町村",
    "種別", "部門", "品目", "産出額（千万円）", "全国順位", "県内順位",
)
STATUSES = {"x", "-", "…"}
YEAR = 2024


def acquire(refresh: bool) -> bytes:
    if refresh or not RAW.exists():
        response = requests.get(SOURCE, timeout=90)
        response.raise_for_status()
        body = response.content
        if not body.startswith(b"PK") or len(body) < 1_000_000:
            raise ValueError("Unexpected source workbook")
        RAW.parent.mkdir(parents=True, exist_ok=True)
        if RAW.exists():
            previous = RAW.read_bytes()
            if previous != body:
                archive = RAW.parent / "archive" / (hashlib.sha256(previous).hexdigest() + ".xlsx")
                archive.parent.mkdir(parents=True, exist_ok=True)
                if not archive.exists():
                    archive.write_bytes(previous)
        RAW.write_bytes(body)
    return RAW.read_bytes()


def build(refresh: bool = False) -> dict:
    source_bytes = acquire(refresh)
    workbook = load_workbook(RAW, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise ValueError("Unexpected workbook sheets")
    rows = workbook.active.iter_rows(values_only=True)
    if tuple(next(rows)) != HEADERS:
        raise ValueError("Source column schema changed")
    products = []
    product_index = {}
    records = {}
    counts = Counter()
    row_count = 0
    for row in rows:
        if not any(v is not None for v in row):
            continue
        if len(row) != 12 or row[0] != YEAR:
            raise ValueError(f"Unexpected year or row: {row[:4]}")
        year, pref_code, city_code, _, prefecture, city, kind, group, item, amount, *_ = row
        code = f"{int(pref_code):02d}{int(city_code):03d}"
        if item not in product_index:
            product_index[item] = len(products)
            products.append({"name": item, "group": group, "kind": kind,
                             "aggregate": item.endswith("計") or item == "農業産出額"})
        idx = product_index[item]
        if products[idx]["group"] != group or products[idx]["kind"] != kind:
            raise ValueError(f"Conflicting item metadata: {item}")
        key = (code, prefecture, city)
        if key not in records:
            records[key] = {"code": code, "prefecture": prefecture, "city": city,
                            "values": {}}
        elif idx in records[key]["values"]:
            raise ValueError(f"Duplicate city/product {key} {item}")
        if isinstance(amount, bool) or (not isinstance(amount, int) and amount not in STATUSES):
            raise ValueError(f"Unexpected amount {key} {item}: {amount!r}")
        if isinstance(amount, int) and amount < 0:
            raise ValueError(f"Negative value: {key} {item}")
        records[key]["values"][idx] = amount
        counts[str(amount) if isinstance(amount, str) else "number"] += 1
        row_count += 1
    if row_count != 120330 or len(records) != 1719 or len(products) != 70:
        raise ValueError(f"Source coverage changed: {row_count} rows, {len(records)} municipalities, {len(products)} products")
    if len({r["code"] for r in records.values()}) != len(records):
        raise ValueError("Duplicate municipality codes")
    if len({r["prefecture"] for r in records.values()}) != 47:
        raise ValueError("Expected 47 prefectures")
    result = []
    for r in records.values():
        if len(r["values"]) != len(products):
            raise ValueError(f"Incomplete municipality: {r['code']}")
        r["values"] = [r["values"][i] for i in range(len(products))]
        result.append(r)
    total_idx = product_index["農業産出額"]
    for code, city, expected in (("32202", "浜田市", 717), ("32204", "益田市", 1019)):
        hit = next((r for r in result if r["code"] == code and r["city"] == city), None)
        if hit is None or hit["values"][total_idx] != expected:
            raise ValueError(f"Corrected September 2026 value absent: {code}")
    data = {
        "schemaVersion": 1, "year": YEAR, "unit": "千万円",
        "source": "農林水産省 市町村別農業産出額（推計） 詳細品目別データ",
        "sourceUrl": SOURCE,
        "sourcePage": "https://www.maff.go.jp/j/tokei/kouhyou/sityoson_sansyutu/",
        "sourceSha256": hashlib.sha256(source_bytes).hexdigest(),
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "municipalityCount": len(result), "sourceRows": row_count,
        "statusCounts": dict(counts), "products": products, "records": result,
    }
    PUBLIC.parent.mkdir(parents=True, exist_ok=True)
    temporary = PUBLIC.with_suffix(".tmp")
    temporary.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    temporary.replace(PUBLIC)
    print(f"{PUBLIC}: {len(result)} cities, {len(products)} items, {row_count} source rows; sha256 {data['sourceSha256']}")
    return data


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--refresh", action="store_true", help="Fetch current official workbook")
    args = parser.parse_args()
    with SourceRun("agriculture_output", "農林水産省 市町村別農業産出額（推計）") as run:
        result = build(args.refresh)
        run.set_metrics(records=result["municipalityCount"], sourceRows=result["sourceRows"], products=len(result["products"]), sourceYear=YEAR)
