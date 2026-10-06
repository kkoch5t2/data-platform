"""Transport audit; all public values, units, identities and cached originals."""
import argparse
import hashlib
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'public/data/transport'
RAW = ROOT / 'data/raw/transport'


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--require-raw', action='store_true')
    args = parser.parse_args()
    meta, stations, commute, usage = [json.loads((OUT/(n+'.json')).read_text()) for n in ('index','stations','commute','usage')]
    rows = stations['records']
    assert len(rows) == meta['stationCount'] >= 10000
    assert len({r['id'] for r in rows}) == len(rows)
    assert stations['years'] == list(range(2011,2025))
    seen = set()
    for r in rows:
        assert r['name'] and r['operator'] and r['line']
        assert len(r['values']) == len(r['states']) == 14
        assert 122 <= r['lon'] <= 154 and 20 <= r['lat'] <= 46
        for i,v in enumerate(r['values']):
            assert r['states'][i] in ('ok','missing','private','absent','other-line','unclassified')
            assert (v is not None) == (r['states'][i] == 'ok')
            if v is not None: assert type(v) is int and v >= 0
        assert not seen.intersection(r['sourceFeatures'])
        seen.update(r['sourceFeatures'])
    assert len(seen) == 10534 and seen == set(range(10534))
    assert sum(r['values'][-1] is not None for r in rows) == meta['stationAvailableCount']
    shinjuku = [r for r in rows if r['name']=='新宿' and r['operator']=='東日本旅客鉄道' and r['values'][-1] is not None]
    assert len(shinjuku) == 1 and shinjuku[0]['values'][-1] == 1333618
    gaps = [r for r in rows if 'unclassified' in r['states']]
    assert len(gaps) == 1 and gaps[0]['code'] == '010175'
    assert [i+2011 for i,s in enumerate(gaps[0]['states']) if s=='unclassified'] == list(range(2012,2019))
    areas=commute['records'];modes=commute['modes']
    assert commute['year']==2020 and len(areas)==meta['commuteAreaCount']==1965
    assert len({a['code'] for a in areas})==len(areas)
    assert len(modes)==17 and len({m['code'] for m in modes})==17
    for a in areas:
        assert re.fullmatch(r'\d{5}',a['code']) and a['name'] and a['prefecture']
        assert len(a['values'])==17 and all(type(v) is int and v>=0 for v in a['values'])
        assert sum(a['values'])==a['total']
    nation=next(a for a in areas if a['code']=='00000')
    assert nation['total']==57152761 and nation['values'][0]==3999367 and nation['values'][1]==9784717
    bus=usage['bus']; assert bus['year']==meta['busYear'] and len(bus['records'])==48
    assert {r['code'] for r in bus['records']}=={str(n).zfill(2) for n in range(48)}
    assert all(type(r['thousands']) is int and r['thousands']>=0 for r in bus['records'])
    assert abs(sum(r['thousands'] for r in bus['records'][1:])-bus['records'][0]['thousands'])<=54
    for kind in ('annual','monthly'):
        rr=usage['rail'][kind]; assert len(rr)>5
        assert [r['period'] for r in rr]==sorted({r['period'] for r in rr})
        for r in rr:
            assert len(r['thousands'])==3 and all(type(v) is int and v>=0 for v in r['thousands'])
            assert abs(r['thousands'][0]-sum(r['thousands'][1:]))<=1
    monthly=usage['rail']['monthly']
    assert monthly[-1]['period']==meta['railLatestMonth']
    indices=[int(r['period'][:4])*12+int(r['period'][5:]) for r in monthly]
    assert indices==list(range(indices[0],indices[-1]+1))
    originals=[RAW/s['file'] for s in meta['sources'].values()]
    if not all(p.exists() for p in originals):
        assert not args.require_raw, 'Original files required for source comparison'
        print('Public structure and published anchors passed; raw originals not present in this checkout.')
        return
    for s in meta['sources'].values():
        payload=(RAW/s['file']).read_bytes()
        assert len(payload)==s['bytes'] and hashlib.sha256(payload).hexdigest()==s['sha256']
    with zipfile.ZipFile(originals[0]) as z:
        features=json.loads(z.read('S12-25_GML/UTF-8/S12-25_NumberOfPassengers.geojson'))['features']
    for r in rows:
        for i in range(14):
            candidates=[]
            for pos in r['sourceFeatures']:
                p=features[pos]['properties']
                assert [r['code'],r['name'],r['operator'],r['line']]==[p[k] for k in ('S12_001c','S12_001','S12_002','S12_003')]
                base=6+i*4
                if p[f'S12_{base:03}']==1 and p[f'S12_{base+1:03}']==1:
                    candidates.append(p[f'S12_{base+3:03}'])
            assert len(set(candidates))<=1
            assert r['values'][i]==(candidates[0] if candidates else None)
    bycode={a['code']:a for a in areas};by_mode={m['code']:i for i,m in enumerate(modes)}
    sheet=openpyxl.load_workbook(originals[1],read_only=True,data_only=True)['e17_02']
    comparisons=0
    for row in sheet.iter_rows(min_row=11,values_only=True):
        if not row[2] or not row[3]: continue
        code=row[2].split('_')[0];mode=row[3].split('_')[0]
        if mode not in by_mode: continue
        value=0 if row[4]=='-' else row[4]
        assert bycode[code]['values'][by_mode[mode]]==value
        comparisons+=1
    assert comparisons==1965*17
    # Independent column anchors plus complete regional counts from the original sheet.
    sheet=openpyxl.load_workbook(originals[2],read_only=True,data_only=True).worksheets[0]
    expected={}
    for row in sheet.iter_rows(min_row=9,max_row=62,values_only=True):
        if row[0]=='全国計': expected['全国']=row[3]
        elif row[1]: expected[row[1]]=row[3]
    for r in bus['records']:
        if r['code']=='01': assert r['thousands']==sum(expected[b['name']] for b in r['branches'])
        else: assert r['thousands']==expected[r['name'].removesuffix('都').removesuffix('府').removesuffix('県')]
    sheet=openpyxl.load_workbook(originals[3],read_only=True,data_only=True)['旅客数量（機械判読用）']
    rr={r['period']:r for kind in ('annual','monthly') for r in usage['rail'][kind]}
    for row in sheet.values:
        label=str(row[0] or '');m=re.fullmatch(r'(20\d{2})年(\d+)月',label);y=re.fullmatch(r'(20\d{2})年度',label)
        if m or y:
            key=y[1] if y else m[1]+'-'+m[2].zfill(2)
            assert rr[key]['thousands']==[row[1],row[3],row[5]]
    print(f'Transport audit passed: {len(rows)} station records / 147476 raw year cells, {comparisons} census cells, 48 bus regions, rail annual/monthly originals.')


if __name__=='__main__': main()
